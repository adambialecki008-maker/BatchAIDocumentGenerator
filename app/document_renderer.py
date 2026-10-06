from docx.document import Document as DocumentObject
from docx import Document
from pathlib import Path
from app.models import ClientInput, GeneratedContent
from datetime import date
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocxRenderer:
    def __init__(
        self,
        resume_template_path: Path,
        cover_letter_template_path: Path,
    ):
        self.resume_template_path = resume_template_path
        self.cover_letter_template_path = cover_letter_template_path

    def render_resume(
        self,
        client: ClientInput,
        content: GeneratedContent,
        output_path: Path,
    ) -> None:
        document = Document(self.resume_template_path)
        replace_placeholder(
            document,
            "{{FULL_NAME}}",
            f"{client.first_name} {client.last_name}",
        )
        replace_placeholder(
            document,
            "{{TARGET_ROLE}}",
            client.target_role,
        )
        replace_placeholder(
            document,
            "{{EMAIL}}",
            client.email,
        )
        replace_placeholder(
            document,
            "{{LOCATION}}",
            client.location,
        )
        replace_placeholder(
            document,
            "{{PROFILE_SUMMARY}}",
            content.professional_summary,
        )
        replace_list_placeholder(
            document,
            "{{SKILLS_BULLETS}}",
            content.key_strengths,
        )
        replace_placeholder(
            document,
            "{{EXPERIENCE_SECTION}}",
            f"{client.current_role} — {client.current_company}",
        )
        replace_placeholder(
            document,
            "{{KEY_ACHIEVEMENT_BULLET}}",
            client.key_achievement,
        )
        style_resume(document)
        document.save(output_path)

    def render_cover_letter(
        self,
        client: ClientInput,
        content: GeneratedContent,
        output_path: Path,
    ) -> None:
        document = Document(self.cover_letter_template_path)
        replace_placeholder(
            document,
            "{{FULL_NAME}}",
            f"{client.first_name} {client.last_name}",
        )
        replace_placeholder(
            document,
            "{{TARGET_COMPANY}}",
            client.target_company,
        )
        replace_placeholder(
            document,
            "{{EMAIL}}",
            client.email,
        )
        replace_placeholder(
            document,
            "{{LOCATION}}",
            client.location,
        )
        replace_placeholder(
            document,
            "{{OPENING_PARAGRAPH}}",
            content.opening_paragraph,
        )
        replace_placeholder(
            document,
            "{{FIT_PARAGRAPH}}",
            content.fit_paragraph,
        )
        replace_placeholder(
            document,
            "{{ACHIEVEMENT_PARAGRAPH}}",
            content.achievement_paragraph,
        )
        replace_placeholder(
            document,
            "{{CLOSING_PARAGRAPH}}",
            content.closing_paragraph,
        )
        replace_placeholder(
            document,
            "{{DATE}}",
            date.today().isoformat(),
        )
        style_cover_letter(document)
        document.save(output_path)


def replace_placeholder(
    document: DocumentObject,
    placeholder: str,
    value: str,
) -> None:
    for paragraph in document.paragraphs:
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(
                placeholder,
                value,
            )


def replace_list_placeholder(
    document: DocumentObject,
    placeholder: str,
    values: list[str],
) -> None:
    text = "\n".join(f"• {value}" for value in values)
    replace_placeholder(
        document,
        placeholder,
        text,
    )


def style_resume(document: DocumentObject) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.05
        text = paragraph.text.strip()
        if text in {
            "PROFILE",
            "CORE SKILLS",
            "EXPERIENCE",
            "KEY ACHIEVEMENT",
        }:
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(11)
    if document.paragraphs:
        name_paragraph = document.paragraphs[0]
        name_paragraph.paragraph_format.space_after = Pt(2)
        for run in name_paragraph.runs:
            run.bold = True
            run.font.size = Pt(18)
    if len(document.paragraphs) > 1:
        role_paragraph = document.paragraphs[1]
        role_paragraph.paragraph_format.space_after = Pt(2)
        for run in role_paragraph.runs:
            run.bold = True
            run.font.size = Pt(12)
    if len(document.paragraphs) > 2:
        contact_paragraph = document.paragraphs[2]
        contact_paragraph.paragraph_format.space_after = Pt(10)
        for run in contact_paragraph.runs:
            run.font.size = Pt(9)


def style_cover_letter(document: DocumentObject) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.space_after = Pt(8)
    if document.paragraphs:
        name_paragraph = document.paragraphs[0]
        name_paragraph.paragraph_format.space_after = Pt(2)
        for run in name_paragraph.runs:
            run.bold = True
            run.font.size = Pt(15)
    if len(document.paragraphs) > 1:
        contact_paragraph = document.paragraphs[1]
        for run in contact_paragraph.runs:
            run.font.size = Pt(9)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "Dear Hiring Manager,":
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(8)

        if text == "Sincerely,":
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(2)
