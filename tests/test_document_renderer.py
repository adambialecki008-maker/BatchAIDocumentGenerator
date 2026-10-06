from datetime import date
from docx import Document
from app.document_renderer import DocxRenderer
from app.models import ClientInput, GeneratedContent


def valid_client_data() -> dict:
    return {
        "client_id": "C001",
        "first_name": "Adam",
        "last_name": "Kowalski",
        "email": "adam@example.com",
        "target_role": "Automation Engineer",
        "years_experience": 5,
        "skills": "PLC, Python",
        "current_company": "ABC Automation",
        "current_role": "Automation Engineer",
        "location": "Kraków",
        "key_achievement": "Commissioned production line",
        "tone": "professional",
        "target_company": "Target Company",
    }


def generated_content() -> GeneratedContent:
    return GeneratedContent(
        professional_summary="Experienced automation engineer.",
        key_strengths=["PLC", "Python"],
        opening_paragraph="I am applying for the position.",
        fit_paragraph="My experience matches the role.",
        achievement_paragraph="I commissioned a production line.",
        closing_paragraph="I would welcome the opportunity to discuss the role.",
    )


def test_render_resume_replaces_client_template_placeholders(tmp_path):
    resume_template_path = tmp_path / "resume_template.docx"
    cover_template_path = tmp_path / "cover_letter_template.docx"
    output_path = tmp_path / "resume.docx"
    template = Document()
    template.add_paragraph("{{FULL_NAME}}")
    template.add_paragraph("{{TARGET_ROLE}}")
    template.add_paragraph("{{EMAIL}} | {{LOCATION}}")
    template.add_paragraph("{{PROFILE_SUMMARY}}")
    template.add_paragraph("{{SKILLS_BULLETS}}")
    template.add_paragraph("{{EXPERIENCE_SECTION}}")
    template.add_paragraph("{{KEY_ACHIEVEMENT_BULLET}}")
    template.save(resume_template_path)
    Document().save(cover_template_path)
    renderer = DocxRenderer(
        resume_template_path,
        cover_template_path,
    )

    renderer.render_resume(
        ClientInput(**valid_client_data()),
        generated_content(),
        output_path,
    )
    assert output_path.exists()
    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Adam Kowalski" in text
    assert "Automation Engineer" in text
    assert "adam@example.com" in text
    assert "Kraków" in text
    assert "Experienced automation engineer." in text
    assert "• PLC" in text
    assert "• Python" in text
    assert "Automation Engineer — ABC Automation" in text
    assert "Commissioned production line" in text
    assert "{{" not in text
    assert "}}" not in text


def test_render_cover_letter_replaces_client_template_placeholders(tmp_path):
    resume_template_path = tmp_path / "resume_template.docx"
    cover_template_path = tmp_path / "cover_letter_template.docx"
    output_path = tmp_path / "cover_letter.docx"
    Document().save(resume_template_path)
    template = Document()
    template.add_paragraph("{{FULL_NAME}}")
    template.add_paragraph("{{EMAIL}} | {{LOCATION}}")
    template.add_paragraph("{{DATE}}")
    template.add_paragraph("{{TARGET_COMPANY}}")
    template.add_paragraph("{{OPENING_PARAGRAPH}}")
    template.add_paragraph("{{FIT_PARAGRAPH}}")
    template.add_paragraph("{{ACHIEVEMENT_PARAGRAPH}}")
    template.add_paragraph("{{CLOSING_PARAGRAPH}}")
    template.add_paragraph("{{FULL_NAME}}")
    template.save(cover_template_path)
    renderer = DocxRenderer(
        resume_template_path,
        cover_template_path,
    )

    renderer.render_cover_letter(
        ClientInput(**valid_client_data()),
        generated_content(),
        output_path,
    )
    assert output_path.exists()
    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Adam Kowalski" in text
    assert "adam@example.com" in text
    assert "Kraków" in text
    assert date.today().isoformat() in text
    assert "Target Company" in text
    assert "I am applying for the position." in text
    assert "My experience matches the role." in text
    assert "I commissioned a production line." in text
    assert "I would welcome the opportunity to discuss the role." in text
    assert "{{" not in text
    assert "}}" not in text
